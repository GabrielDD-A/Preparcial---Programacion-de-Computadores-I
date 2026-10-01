# -*- coding: utf-8 -*-
"""Preparcial Programacion de computadores I"""

import pandas as pd
import ipywidgets as widgets
from IPython.display import display, clear_output

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

try:
    from google.colab import files
    enColab = True
except ImportError:
    enColab = False


def formatoCop(valor, decimales=0):
    texto = f"{valor:,.{decimales}f}"
    texto = texto.replace(",", "X").replace(".", ",").replace("X", ".")
    return "$" + texto


turno = {
    "cajero": None,
    "tope": None,
    "recaudoTotal": 0,
    "transacciones": 0,
    "activa": False,
    "motivoCierre": None
}

detalleActual = []
detalleTransacciones = []
reportes = []

nombreExcel = "caja_recaudo.xlsx"
nombreCsvReportes = "reportes_turnos.csv"
nombreCsvDetalle = "detalle_transacciones.csv"

columnasMonedaEnteras = ["Tope", "Recaudo_total", "Monto", "Acumulado"]
columnasMonedaDecimales = ["Promedio_transaccion"]

bordeFino = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0")
)


def aplicarFormatoHoja(nombreArchivo, nombreHoja, df):
    wb = load_workbook(nombreArchivo)
    ws = wb[nombreHoja]

    if df.empty:
        wb.save(nombreArchivo)
        return

    nFilas = len(df) + 1
    nColumnas = len(df.columns)

    rellenoEncabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col in range(1, nColumnas + 1):
        celda = ws.cell(row=1, column=col)
        celda.fill = rellenoEncabezado
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center")
        celda.border = bordeFino

    for col in range(1, nColumnas + 1):
        letra = get_column_letter(col)
        largoMax = max(df.iloc[:, col - 1].astype(str).map(len).max(), len(str(df.columns[col - 1])))
        ws.column_dimensions[letra].width = largoMax + 4

    for col in range(1, nColumnas + 1):
        nombreCol = df.columns[col - 1]
        for fila in range(2, nFilas + 1):
            celda = ws.cell(row=fila, column=col)
            if nombreCol in columnasMonedaEnteras:
                celda.number_format = '$ #,##0'
            elif nombreCol in columnasMonedaDecimales:
                celda.number_format = '$ #,##0.00'
            celda.alignment = Alignment(horizontal="center")
            celda.border = bordeFino

    rango = f"A1:{get_column_letter(nColumnas)}{nFilas}"
    tabla = Table(displayName=f"Tabla_{nombreHoja}", ref=rango)
    tabla.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tabla)

    ws.freeze_panes = "A2"
    wb.save(nombreArchivo)


def guardarExcelAutomatico():
    columnasReporte = ["Cajero", "Tope", "Transacciones", "Recaudo_total",
                       "Promedio_transaccion", "Motivo_cierre"]
    dfReportes = pd.DataFrame(reportes, columns=columnasReporte)

    columnasDetalle = ["Cajero", "Num_transaccion", "Monto", "Acumulado"]
    dfDetalle = pd.DataFrame(detalleTransacciones, columns=columnasDetalle)

    with pd.ExcelWriter(nombreExcel, engine="openpyxl") as writer:
        dfReportes.to_excel(writer, sheet_name="Reportes", index=False)
        dfDetalle.to_excel(writer, sheet_name="Detalle", index=False)

    aplicarFormatoHoja(nombreExcel, "Reportes", dfReportes)
    aplicarFormatoHoja(nombreExcel, "Detalle", dfDetalle)

    dfReportes.to_csv(nombreCsvReportes, index=False, sep=";", encoding="utf-8-sig")
    dfDetalle.to_csv(nombreCsvDetalle, index=False, sep=";", encoding="utf-8-sig")


cajeroInput = widgets.Text(description="Cajero:", placeholder="Ej: C-102")
topeInput = widgets.FloatText(description="Tope ($):", value=0)
botonIniciar = widgets.Button(description="Iniciar turno", button_style="primary")

montoInput = widgets.Text(description="Monto:", placeholder="Escriba el monto, o 0 / FIN para cerrar")
botonRegistrar = widgets.Button(description="Registrar transaccion", button_style="success")
botonFinalizar = widgets.Button(description="Finalizar turno (FIN)", button_style="warning")

botonCsv = widgets.Button(description="Exportar CSV")
botonExcel = widgets.Button(description="Exportar Excel")

salida = widgets.Output()


def iniciarTurno(boton):
    global detalleActual

    if not cajeroInput.value.strip():
        with salida:
            clear_output()
            print("Debe indicar el nombre o codigo del cajero.")
        return

    if topeInput.value <= 0:
        with salida:
            clear_output()
            print("El tope debe ser un numero mayor que cero.")
        return

    turno["cajero"] = cajeroInput.value.strip()
    turno["tope"] = topeInput.value
    turno["recaudoTotal"] = 0
    turno["transacciones"] = 0
    turno["activa"] = True
    turno["motivoCierre"] = None
    detalleActual = []

    with salida:
        clear_output()
        print(f"Turno iniciado para el cajero {turno['cajero']}.")
        print(f"Tope asignado: {formatoCop(turno['tope'])}")

botonIniciar.on_click(iniciarTurno)


def cerrarTurno(motivo):
    turno["activa"] = False
    turno["motivoCierre"] = motivo

    if turno["transacciones"] > 0:
        promedio = turno["recaudoTotal"] / turno["transacciones"]
    else:
        promedio = 0

    reportes.append([
        turno["cajero"], turno["tope"], turno["transacciones"],
        turno["recaudoTotal"], promedio, turno["motivoCierre"]
    ])

    guardarExcelAutomatico()

    with salida:
        clear_output()
        print("===== REPORTE DE TURNO =====")
        print(f"Cajero:               {turno['cajero']}")
        print(f"Tope asignado:        {formatoCop(turno['tope'])}")
        print(f"Transacciones:        {turno['transacciones']}")
        print(f"Recaudo total:        {formatoCop(turno['recaudoTotal'])}")
        print(f"Promedio/transaccion: {formatoCop(promedio, 2)}")
        print(f"Motivo de cierre:     {turno['motivoCierre']}")
        print()
        print(f"(Guardado automaticamente en {nombreExcel}, con formato de tabla)")


def registrar(boton):
    if not turno["activa"]:
        with salida:
            clear_output()
            print("No hay un turno activo. Inicie turno primero.")
        return

    texto = montoInput.value.strip()

    if texto == "0" or texto.upper() == "FIN":
        cerrarTurno("Fin de cola")
        montoInput.value = ""
        return

    try:
        monto = float(texto)
    except ValueError:
        with salida:
            clear_output()
            print("Monto invalido: debe ser un numero. Intente de nuevo.")
        montoInput.value = ""
        return

    if monto <= 0:
        with salida:
            clear_output()
            print("Monto invalido: debe ser mayor que cero. Intente de nuevo.")
        montoInput.value = ""
        return

    turno["recaudoTotal"] += monto
    turno["transacciones"] += 1

    fila = [turno["cajero"], turno["transacciones"], monto, turno["recaudoTotal"]]
    detalleActual.append(fila)
    detalleTransacciones.append(fila)

    guardarExcelAutomatico()

    with salida:
        clear_output()
        print(f"Transaccion #{turno['transacciones']} registrada: {formatoCop(monto)}")
        print(f"Acumulado: {formatoCop(turno['recaudoTotal'])}")

    if turno["recaudoTotal"] >= turno["tope"]:
        cerrarTurno("Tope alcanzado")
        with salida:
            print()
            print("CAJA SUSPENDIDA: se alcanzo el tope de recaudo. Dirijase a tesoreria.")

    montoInput.value = ""

botonRegistrar.on_click(registrar)


def finalizar(boton):
    if not turno["activa"]:
        with salida:
            clear_output()
            print("No hay un turno activo.")
        return
    cerrarTurno("Fin de cola")

botonFinalizar.on_click(finalizar)


def exportarCsv(boton):
    guardarExcelAutomatico()
    with salida:
        clear_output()
        print("CSV actualizados:", nombreCsvReportes, "y", nombreCsvDetalle)
    if enColab:
        files.download(nombreCsvReportes)
        files.download(nombreCsvDetalle)

def exportarExcel(boton):
    guardarExcelAutomatico()
    with salida:
        clear_output()
        print("Excel actualizado y con formato de tabla:", nombreExcel)
    if enColab:
        files.download(nombreExcel)

botonCsv.on_click(exportarCsv)
botonExcel.on_click(exportarExcel)


display(widgets.HTML("<h2>Caja de Recaudo con Tope de Seguridad</h2>"))

display(widgets.HTML("<h3>Inicio de turno</h3>"))
display(cajeroInput)
display(topeInput)
display(botonIniciar)

display(widgets.HTML("<h3>Atencion de clientes</h3>"))
display(montoInput)
display(botonRegistrar)
display(botonFinalizar)

display(widgets.HTML("<h3>Exportar informacion</h3>"))
display(botonCsv)
display(botonExcel)

display(salida)
